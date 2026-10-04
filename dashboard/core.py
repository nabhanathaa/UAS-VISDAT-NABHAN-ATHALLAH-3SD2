# -*- coding: utf-8 -*-
"""Konstanta, palet, utilitas grafik, dan pemuatan data (dipakai semua tab)."""
import base64
import json
from pathlib import Path

import numpy as np
import pandas as pd
import streamlit as st

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "data" / "processed"

# ============================================================== warna & tipografi
NAVY, NAVY2, ROYAL, GOLD = "#0a1f3d", "#123a6b", "#1d5fae", "#f2b544"
INK, INK2, MUTED, GRID, LINE = "#13233a", "#45566d", "#7a8799", "#e8edf4", "#dbe3ee"
FONT = "Plus Jakarta Sans, Inter, Segoe UI, system-ui, sans-serif"
CAT4 = ["#2a78d6", "#eb6834", "#1baf7a", "#4a3aa7"]                      # lolos validator CVD (all-pairs)
SEQ5 = ["#cde2fb", "#86b6ef", "#3987e5", "#1c5cab", "#0d366b"]           # sekuensial biru satu-hue
SEQ_CONT = [[0, "#e6f0fc"], [.25, "#9ec5f4"], [.5, "#3987e5"], [.75, "#1c5cab"], [1, "#0d366b"]]
DIV = [[0, "#1c5cab"], [.25, "#86b6ef"], [.5, "#f0efec"], [.75, "#ee8a88"], [1, "#b42e2e"]]
RED, BLUE_D = "#b42e2e", "#1c5cab"
GRAY = "#c9c8c3"
PULAU = ["Sumatera", "Jawa", "Bali-Nusa Tenggara", "Kalimantan", "Sulawesi", "Maluku", "Papua"]
PULAU_C = dict(zip(PULAU, ["#2a78d6", "#eb6834", "#1baf7a", "#eda100", "#e87ba4", "#008300", "#4a3aa7"]))  # 7 slot lolos CVD
KUADRAN_C = {"Terang – Sejahtera": CAT4[0], "Gelap – Sejahtera": CAT4[2],
             "Terang – Kurang sejahtera": CAT4[1], "Gelap – Kurang sejahtera": CAT4[3]}
STATUS_C = {"Selaras": "#c9c8c3", "Lebih kaya dari cahayanya": RED, "Lebih terang dari ekonominya": BLUE_D}
LISA_C = {"HH (kaya-dari-cahaya mengelompok)": RED, "LL (terang-dari-ekonomi mengelompok)": BLUE_D,
          "HL": "#ee8a88", "LH": "#86b6ef", "Tidak signifikan": "#e8e7e3"}
KLASTER_C = {"Perkotaan padat & terang": CAT4[0], "Perdesaan berkembang": CAT4[1],
             "Kantong pertumbuhan pesat (SDA/IKN)": CAT4[2], "Pedalaman tertinggal: gelap & mahal": CAT4[3]}
CAT_MAP = {"kuadran": KUADRAN_C, "klaster": KLASTER_C, "status_selaras": STATUS_C, "lisa": LISA_C}
SRC = "Sumber: BPS (2024) · VIIRS VNL V2.1/V2.2, EOG via Google Earth Engine · batas wilayah: geoBoundaries (BPS–OCHA)"
GEO_KAB, GEO_PROV = "app/static/kabkota.geojson", "app/static/provinsi.geojson"   # disajikan statis & di-cache peramban

# variabel: kolom → (label, satuan, jenis skala, format, tersedia di provinsi, kelompok)
VAR = {
    "cahaya_pk":          ("Cahaya malam per kapita", "nW/cm²/sr per 1.000 jiwa", "seq", ".2f", True, "Cahaya"),
    "pct_menyala":        ("Luas wilayah menyala", "% luas", "seq", ".1f", True, "Cahaya"),
    "pertumbuhan_cahaya": ("Pertumbuhan cahaya 2020–2024", "% (log-beda)", "div", ".1f", True, "Cahaya"),
    "pdrb_kapita_juta":   ("PDRB per kapita ADHB", "juta Rp", "seq", ".1f", True, "Ekonomi"),
    "laju_pdrb_pct":      ("Laju pertumbuhan PDRB", "%", "div", ".2f", True, "Ekonomi"),
    "ikk":                ("Indeks Kemahalan Konstruksi", "indeks (nasional = 100)", "seq", ".1f", True, "Ekonomi"),
    "ipm":                ("Indeks Pembangunan Manusia", "indeks", "seq", ".2f", True, "Kesejahteraan"),
    "p0":                 ("Penduduk miskin (P0)", "% penduduk", "seq", ".2f", True, "Kesejahteraan"),
    "kepadatan":          ("Kepadatan penduduk", "jiwa/km²", "seq", ",.0f", True, "Kesejahteraan"),
    "z_residual":         ("Ketidakselarasan cahaya–ekonomi", "z residual OLS", "div", ".2f", False, "Analisis"),
    "kuadran":            ("Kuadran cahaya × kesejahteraan", "", "cat", "", False, "Analisis"),
    "status_selaras":     ("Status keselarasan", "", "cat", "", False, "Analisis"),
    "lisa":               ("Klaster spasial LISA (residual)", "", "cat", "", False, "Analisis"),
    "klaster":            ("Tipologi k-means", "", "cat", "", False, "Analisis"),
}
PENDEK = {"Cahaya per kapita (ln)": "Cahaya/kap", "% luas menyala": "% menyala", "Pertumbuhan cahaya 2020–24": "Δ cahaya",
          "PDRB per kapita (ln)": "PDRB/kap", "Laju PDRB 2024": "Laju PDRB", "IPM": "IPM", "Penduduk miskin (P0)": "P0",
          "Indeks Kemahalan Konstruksi": "IKK", "Kepadatan (ln)": "Kepadatan"}
KAT_PENDEK = {"A": "Pertanian", "B": "Pertambangan", "C": "Industri pengolahan", "D": "Listrik & gas", "E": "Air & sampah",
              "F": "Konstruksi", "G": "Perdagangan", "H": "Transportasi", "I": "Akomodasi & makan", "J": "Infokom",
              "K": "Keuangan", "L": "Real estat", "M,N": "Jasa perusahaan", "O": "Adm. pemerintahan",
              "P": "Pendidikan", "Q": "Kesehatan", "R,S,T,U": "Jasa lainnya"}
KAT_C = dict(zip(["Industri pengolahan", "Perdagangan", "Pertanian", "Konstruksi", "Pertambangan", "Transportasi", "Infokom"],
                 ["#2a78d6", "#eb6834", "#008300", "#eda100", "#4a3aa7", "#1baf7a", "#e87ba4"]))
CHI2_975_9 = 19.02   # batas χ²(0,975; df = 9) untuk pencilan Mahalanobis


# ============================================================== utilitas
def angka(x, d=2, tanda=False):
    """Format angka gaya Indonesia: 1.234,56"""
    try:
        if x is None or np.isnan(float(x)):
            return "–"
    except (TypeError, ValueError):
        return str(x)
    t = f"{x:+,.{d}f}" if tanda else f"{x:,.{d}f}"
    return t.replace(",", "_").replace(".", ",").replace("_", ".")


def fmtv(x, fmt):
    d = 0 if fmt.startswith(",.0") else int(fmt[-2]) if len(fmt) > 1 and fmt[-2].isdigit() else 2
    return angka(x, d)


def pendek(n):
    return str(n).replace("Kabupaten ", "Kab. ")


def pusat(D, W=900, H=560, pad=.6):
    """Pusat & zoom otomatis mengikuti wilayah terfilter (zoom-to-extent; kuantil 1–99% agar pulau terluar tak mendominasi)."""
    q = (0.01, 0.99) if len(D) > 40 else (0, 1)
    lo0, lo1 = D.lon.quantile(q[0]) - pad, D.lon.quantile(q[1]) + pad
    la0, la1 = D.lat.quantile(q[0]) - pad, D.lat.quantile(q[1]) + pad
    z = min(np.log2(W * 360 / (512 * (lo1 - lo0))), np.log2(H * 360 / (512 * (la1 - la0))))
    return dict(center=dict(lat=float((la0 + la1) / 2), lon=float((lo0 + lo1) / 2)), zoom=float(np.clip(z - .1, 2, 7.5)))


def tema(fig, legend=False, l=8, r=8, t=8, b=8):
    fig.update_layout(height=None, autosize=True, margin=dict(l=l, r=r, t=t, b=b), separators=",.",
                      paper_bgcolor="rgba(0,0,0,0)", plot_bgcolor="rgba(0,0,0,0)",
                      font=dict(family=FONT, size=11.5, color=INK2),
                      hoverlabel=dict(bgcolor="white", bordercolor=LINE, font=dict(family=FONT, size=12, color=INK)),
                      showlegend=legend, legend=dict(font=dict(size=11), bgcolor="rgba(255,255,255,.85)", title=None))
    fig.update_xaxes(gridcolor=GRID, zeroline=False, linecolor=LINE, title_font=dict(size=11.5, color=INK2), tickfont=dict(size=10.5))
    fig.update_yaxes(gridcolor=GRID, zeroline=False, linecolor=LINE, title_font=dict(size=11.5, color=INK2), tickfont=dict(size=10.5))
    return fig


def step_cs(colors):
    n = len(colors)
    cs = []
    for i, c in enumerate(colors):
        cs += [[i / n, c], [(i + 1) / n, c]]
    return cs


def cbar(title="", **kw):
    d = dict(thickness=11, len=.6, x=.99, xanchor="right", y=.03, yanchor="bottom", bgcolor="rgba(255,255,255,.9)",
             outlinewidth=0, tickfont=dict(size=10.5, color=INK), title=dict(text=title, font=dict(size=10.5, color=INK2), side="top"))
    d.update(kw)
    return d


def tampil(fig, key, cfg=None, **kw):
    conf = {"displaylogo": False, "responsive": True, "scrollZoom": True,
            "modeBarButtonsToRemove": ["select2d", "lasso2d", "toImage", "autoScale2d"] if not kw.get("on_select") else ["toImage"]}
    conf.update(cfg or {})
    return st.plotly_chart(fig, key=key, height="stretch", config=conf, **kw)


def kepala(judul, sub="", chips=()):
    ch = "".join(f"<span class='chip'>{c}</span>" for c in chips)
    st.html(f"<div class='kh'><div class='kt'>{judul}</div><div class='chips'>{ch}</div></div>"
            + (f"<div class='ks'>{sub}</div>" if sub else ""))


def sumber(teks=SRC):
    st.html(f"<div class='src'>{teks}</div>")


def kpi_grid(items):
    """items: list of (label, nilai, keterangan, warna aksen)."""
    h = "".join(f"<div class='kpi' style='--ac:{c}'><div class='kl'>{a}</div><div class='kv'>{b}</div><div class='kd'>{d}</div></div>"
                for a, b, d, c in items)
    st.html(f"<div class='kpis' style='grid-template-columns:repeat({len(items)},1fr)'>{h}</div>")


@st.cache_data
def logo_b64():
    return base64.b64encode((ROOT / "assets" / "logo_stis_128.png").read_bytes()).decode()


def rgba(h, a):
    return f"rgba({int(h[1:3], 16)},{int(h[3:5], 16)},{int(h[5:7], 16)},{a})"


# ============================================================== data (di-cache sekali per server)
@st.cache_data(show_spinner="Memuat data…")
def muat():
    kab = pd.read_parquet(DATA / "kabkota_analisis.parquet")
    prov = pd.read_parquet(DATA / "provinsi.parquet")
    sek = pd.read_parquet(DATA / "sektor_provinsi.parquet")
    od = pd.read_parquet(DATA / "migrasi_od.parquet")
    neto = pd.read_parquet(DATA / "migrasi_neto.parquet")
    mat = pd.read_parquet(DATA / "migrasi_matriks.parquet")
    ring = json.loads((DATA / "ringkasan_analisis.json").read_text(encoding="utf-8"))
    kab["pdrb_kapita_juta"] = kab.pdrb_kapita_ribu / 1000
    prov["pdrb_kapita_juta"] = prov.pdrb_kapita_ribu / 1000
    # PDRB ADHB provinsi resmi = Σ 17 lapangan usaha (BPS) → ukuran simbol proporsional tingkat provinsi
    prov["pdrb_adhb_miliar"] = prov.kode_prov.map(sek.groupby("kode_prov").adhb_2024_miliar.sum())
    # Location Quotient (spesialisasi sektor): LQ = (X_pk / X_p) / (X_k / X)
    sek = sek.copy()
    sek["sektor"] = sek.kode_kategori.map(KAT_PENDEK)
    X = sek.adhb_2024_miliar.sum()
    tot_p = sek.groupby("provinsi").adhb_2024_miliar.transform("sum")
    x_pk = sek.groupby(["provinsi", "sektor"]).adhb_2024_miliar.transform("sum")
    x_k = sek.groupby("sektor").adhb_2024_miliar.transform("sum")
    sek["lq_sektor"] = (x_pk / tot_p) / (x_k / X)
    sek["lq_sub"] = (sek.adhb_2024_miliar / tot_p) / (sek.groupby("subkategori").adhb_2024_miliar.transform("sum") / X)
    sek["pangsa_sektor"] = x_pk / tot_p * 100
    sek["daun_sub"] = sek.subkategori != sek.kategori
    # Jarak Mahalanobis (pencilan multivariat) pada 9 variabel terstandar
    cols = list(ring["var_label"])
    Z = (kab[cols] - kab[cols].mean()) / kab[cols].std()
    inv = np.linalg.pinv(np.cov(Z.values, rowvar=False))
    kab["mahal"] = np.einsum("ij,jk,ik->i", Z.values, inv, Z.values)
    return kab, prov, sek, od, neto, mat, ring, Z


# ============================================================== unduhan Excel (2 sheet data + keterangan)
KOLOM_KAB = {
    "kode_bps": "Kode BPS", "kode_kemendagri": "Kode Kemendagri", "nama": "Kabupaten/Kota", "jenis": "Jenis", "provinsi": "Provinsi",
    "kode_prov": "Kode Provinsi", "pulau": "Pulau",
    "penduduk_2024": "Penduduk 2024 (jiwa)", "kepadatan": "Kepadatan 2024 (jiwa/km²)", "luas_km2": "Luas poligon (km², dihitung)",
    "pdrb_adhb_miliar": "PDRB ADHB 2024 (miliar Rp)", "pdrb_kapita_ribu": "PDRB per kapita ADHB 2024 (ribu Rp)",
    "laju_pdrb_pct": "Laju pertumbuhan PDRB 2024 (%)", "ipm": "IPM 2024", "p0": "P0 – % penduduk miskin 2024",
    "p1": "P1 – indeks kedalaman kemiskinan", "p2": "P2 – indeks keparahan kemiskinan", "gk": "Garis kemiskinan (Rp/kapita/bulan)",
    "miskin_ribu": "Jumlah penduduk miskin (ribu jiwa)", "ikk": "IKK 2024 (nasional = 100)",
    "rad_sum_2020": "Σ radiance VIIRS 2020 (nW/cm²/sr)", "rad_sum_2024": "Σ radiance VIIRS 2024 (nW/cm²/sr)",
    "rad_mean_2024": "Rata-rata radiance per piksel 2024", "cfcvg_mean_2024": "Rata-rata malam bebas awan 2024",
    "cahaya_pk": "Cahaya per kapita 2024 (nW/cm²/sr per 1.000 jiwa)", "pct_menyala": "Luas menyala 2024 (% piksel > 0,5)",
    "pertumbuhan_cahaya": "Pertumbuhan cahaya 2020–2024 (% log-beda)", "flag_awan": "Mutu awan rendah (<15 malam)",
    "pdrb_prediksi_ribu": "PDRB/kap prediksi OLS (ribu Rp)", "z_residual": "z residual OLS",
    "status_selaras": "Status keselarasan", "indeks_sejahtera": "Indeks kesejahteraan [z(IPM)−z(P0)]/2", "kuadran": "Kuadran",
    "lisa": "Klaster LISA (residual)", "lisa_p": "p-value LISA", "PC1": "Skor PC1", "PC2": "Skor PC2", "PC3": "Skor PC3",
    "klaster": "Tipologi k-means", "mahal": "Jarak Mahalanobis (9 variabel)", "lon": "Bujur titik representatif", "lat": "Lintang titik representatif",
}
KOLOM_PROV = {
    "kode_prov": "Kode Provinsi", "provinsi": "Provinsi", "pulau": "Pulau", "jumlah_kabkota": "Jumlah kab/kota",
    "penduduk_2024": "Penduduk 2024 (jiwa)", "kepadatan": "Kepadatan 2024 (jiwa/km²)",
    "pdrb_adhb_miliar": "PDRB ADHB 2024 = Σ17 lapangan usaha (miliar Rp)", "pdrb_kapita_ribu": "PDRB per kapita ADHB 2024 (ribu Rp)",
    "laju_pdrb_pct": "Laju PDRB 2024 (%; rata-rata tertimbang PDRB kab/kota)", "ipm": "IPM 2024", "p0": "P0 – % penduduk miskin 2024",
    "ikk": "IKK 2024", "rad_sum_2020": "Σ radiance VIIRS 2020", "rad_sum_2024": "Σ radiance VIIRS 2024",
    "cahaya_pk": "Cahaya per kapita 2024 (nW/cm²/sr per 1.000 jiwa)", "pct_menyala": "Luas menyala 2024 (%)",
    "pertumbuhan_cahaya": "Pertumbuhan cahaya 2020–2024 (% log-beda)",
    "migrasi_masuk": "Migrasi masuk risen 2015–2020 (jiwa)", "migrasi_keluar": "Migrasi keluar risen 2015–2020 (jiwa)",
    "migrasi_neto": "Migrasi neto risen 2015–2020 (jiwa)", "lon": "Bujur", "lat": "Lintang",
}
KETERANGAN = [
    ("Judul", "Terang, Kaya, Sejahtera? — data terolah dashboard UAS Visualisasi Data dan Informasi 2026"),
    ("Penyusun", "Nabhan Athallah · 3SD2 · 222313272 · Politeknik Statistika STIS"),
    ("Sheet KabKota", "514 kabupaten/kota: indikator BPS 2024 + cahaya malam VIIRS + hasil analisis (OLS, kuadran, LISA, PCA, k-means)"),
    ("Sheet Provinsi", "38 provinsi: indikator BPS 2024 + agregasi cahaya + migrasi risen LF SP2020 (34 provinsi lama: Papua & Papua Barat = wilayah sebelum pemekaran 2022; 4 provinsi baru kosong)"),
    ("Sumber BPS", "PDRB Kab/Kota 2020–2024; IPM metode baru 2024; Data & Informasi Kemiskinan Kab/Kota 2024; IKK 2024; Jumlah Penduduk menurut Kab/Kota & Kelompok Umur 2024; Kepadatan penduduk 2024; PDRB Provinsi menurut Lapangan Usaha 2021–2025; Statistik Migrasi Indonesia LF SP2020"),
    ("Sumber non-BPS", "NOAA/EOG VIIRS VNL V2.1 (2020) & V2.2 (2024) via Google Earth Engine; batas wilayah geoBoundaries IDN ADM2"),
    ("Cahaya per kapita", "Σ radiance piksel dalam poligon ÷ penduduk × 1.000"),
    ("Luas menyala", "jumlah piksel radiance > 0,5 nW/cm²/sr ÷ jumlah piksel × 100"),
    ("Pertumbuhan cahaya", "100 × [ln(Σrad2024 + 1) − ln(Σrad2020 + 1)]"),
    ("OLS", "ln(PDRB/kap) = a + β·ln(1 + cahaya/kap); galat baku HC3; z residual = residual terstandar; |z| > 1,5 = tidak selaras"),
    ("Kuadran", "terang = ln(1+cahaya/kap) ≥ median; sejahtera = [z(IPM) − z(P0)]/2 ≥ median"),
    ("Tipologi", "k-means (k = 4, silhouette tertinggi) pada skor PC1–PC3 dari 9 variabel terstandar"),
    ("Format angka", "Desimal memakai titik (format mesin). Diakses 2–3 Oktober 2026."),
]


@st.cache_data(show_spinner=False)
def excel_unduhan():
    import io
    kab, prov, _, _, neto, _, _, _ = muat()
    K = kab[[c for c in KOLOM_KAB if c in kab]].rename(columns=KOLOM_KAB)
    neto = neto.replace({"provinsi": {"Kep. Bangka Belitung": "Kepulauan Bangka Belitung"}})
    P = prov.merge(neto[["provinsi", "masuk_risen", "keluar_risen", "neto_risen"]].rename(columns={
        "masuk_risen": "migrasi_masuk", "keluar_risen": "migrasi_keluar", "neto_risen": "migrasi_neto"}),
        on="provinsi", how="left")
    P = P[[c for c in KOLOM_PROV if c in P]].rename(columns=KOLOM_PROV)
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        pd.DataFrame(KETERANGAN, columns=["Butir", "Keterangan"]).to_excel(xw, sheet_name="Keterangan", index=False)
        K.to_excel(xw, sheet_name="KabKota", index=False)
        P.to_excel(xw, sheet_name="Provinsi", index=False)
        from openpyxl.styles import Alignment, Font, PatternFill
        for ws in xw.book.worksheets:
            ws.freeze_panes = "B2" if ws.title != "Keterangan" else "A2"
            for c in ws[1]:
                c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="123A6B")
                c.alignment = Alignment(wrap_text=True, vertical="center")
            ws.row_dimensions[1].height = 42
            for col in ws.columns:
                ws.column_dimensions[col[0].column_letter].width = 16
        xw.book["Keterangan"].column_dimensions["A"].width = 20
        xw.book["Keterangan"].column_dimensions["B"].width = 120
        xw.book["KabKota"].column_dimensions["C"].width = 28
        xw.book["Provinsi"].column_dimensions["B"].width = 26
    return buf.getvalue()
